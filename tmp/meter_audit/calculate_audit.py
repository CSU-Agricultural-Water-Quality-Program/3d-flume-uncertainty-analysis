from pathlib import Path
import csv, hashlib, json, math
from collections import Counter
from openpyxl import load_workbook

ROOT = Path.cwd()
OUT = ROOT / 'output/audit/2026-10-13-flow-meter'
OUT.mkdir(parents=True, exist_ok=True)
K = 28.316846592
input_file = ROOT / 'data/New Data Oct 2026/flume_data_oct2026.csv'
book_file = ROOT / 'data/New Data Oct 2026/Oct 2026 Trapezoidal Flume Calibration Final.xlsx'
rows = list(csv.DictReader(input_file.open(encoding='utf-8-sig')))
wb = load_workbook(book_file, read_only=True, data_only=True)
ws = wb['Detailed Calibration (10-point)']
heads = {'3D1': (4, 'C8'), '3D2': (10, 'I8'), 'FIBER': (16, 'O8')}
points = []
for flume, (col, bottom) in heads.items():
    matching = [r for r in rows if r['flume_type'] == flume]
    assert len(matching) == 10
    for i, r in enumerate(matching, start=10):
        q_cfs = float(r['flow_cfs'])
        h_ft = float(r['depth_ft'])
        assert ws.cell(i, col-2).value == q_cfs
        h_cm_raw = ws.cell(i, col).value - ws[bottom].value
        assert abs(h_ft - h_cm_raw/30.48) <= .0005 + 1e-12
        q = q_cfs*K
        a100 = math.pi*.1**2/4
        a300 = math.pi*.3**2/4
        e_mod = .002*q + a100  # 1 mm/s -> area * 0.001 m/s * 1000 L/m^3
        e_pro = .005*q + a300
        points.append(dict(flume=flume,flow_cfs=q_cfs,flow_lps=q,
                           csv_depth_cm=h_ft*30.48,source_depth_cm=h_cm_raw,
                           stage_rounding_difference_cm=h_ft*30.48-h_cm_raw,
                           modmag_nominal_diameter_m=.1,modmag_velocity_mps=q/1000/a100,
                           modmag_error_bound_lps=e_mod,modmag_error_bound_percent=100*e_mod/q,
                           modmag_rectangular_standard_uncertainty_lps=e_mod/math.sqrt(3),
                           operational_bound_lps=.005*K,
                           operational_rectangular_standard_uncertainty_lps=.005*K/math.sqrt(3),
                           above_modmag_listed_range_min=q>.85/3.6,
                           below_promag_factory_cutoff=q<10/3.6,
                           promag_comparison_error_bound_lps=e_pro))
with (OUT/'meter_error_by_calibration_point.csv').open('w',newline='',encoding='utf-8-sig') as f:
    writer=csv.DictWriter(f,fieldnames=points[0]); writer.writeheader(); writer.writerows(points)
old=list(csv.DictReader((ROOT/'data/flume_data.csv').open(encoding='utf-8-sig')))
old=[r for r in old if r['flume_type']!='EQ']
summary = dict(n=len(points),counts=dict(Counter(p['flume'] for p in points)),
    old_counts=dict(Counter(r['flume_type'] for r in old)),
    old_below_promag_factory_cutoff=[r for r in old if float(r['flow_cfs'])*K<10/3.6],
    lps_per_cfs=K,modmag_listed_range_min_lps=.85/3.6,
    promag_factory_cutoff_lps=10/3.6,
    stage_csv_max_abs_rounding_error_cm=max(abs(p['stage_rounding_difference_cm']) for p in points),
    modmag_min_velocity_mps=min(p['modmag_velocity_mps'] for p in points),
    modmag_max_velocity_mps=max(p['modmag_velocity_mps'] for p in points),
    uncertainty_assumption='Uniform bounds; SD=half-width/sqrt(3). Gaussian likelihood would be an approximation.',
    diameter_assumption='Nominal DN100 and DN300, not verified installed internal diameters.',
    model_modified=False, model_rerun=False)
(OUT/'audit_calculations.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
manifest=json.loads((ROOT/'tmp/meter_audit/source_manifest.json').read_text(encoding='utf-8'))
for p in [input_file,ROOT/'code/bayesflume.r',ROOT/'tmp/meter_audit/manuscript_open_word.xml']:
    manifest.append(dict(name=p.name,path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
(OUT/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
