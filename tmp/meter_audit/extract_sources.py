from pathlib import Path
import hashlib, json, re, zipfile
from pypdf import PdfReader
from lxml import etree
from openpyxl import load_workbook

ROOT = Path("D:/OneDrive - Colostate/Cabot,Perry's files - 3D Printing Flumes")
OUT = Path(__file__).parent
files = {
    'modmag': ROOT / 'ModMag Manual.pdf',
    'modmag_reference': ROOT / 'REFERENCES/MAG-UM-01272-EN_M2000.pdf',
    'promag': ROOT / 'REFERENCES/Promag Manual (1).pdf',
    'oct_workbook': Path('data/New Data Oct 2026/Oct 2026 Trapezoidal Flume Calibration Final.xlsx'),
    'ben_workbook': ROOT / 'Flume Calibration Report/bms 9_17 calibration_BenS.xlsx',
}
if (OUT / 'manuscript_snapshot.docx').exists() and (OUT / 'manuscript_snapshot.docx').stat().st_size:
    files['manuscript'] = OUT / 'manuscript_snapshot.docx'
manifest = []
for name, path in files.items():
    manifest.append(dict(name=name, path=str(path.resolve()), sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    if path.suffix == '.pdf':
        doc = PdfReader(path).pages
        text = '\n'.join(f'\n=== PDF PAGE {i+1} ===\n'+p.extract_text() for i,p in enumerate(doc))
        (OUT / f'{name}.txt').write_text(text, encoding='utf-8')
        print(name, 'pages', len(doc))
        for i,p in enumerate(doc):
            t = p.extract_text()
            if re.search(r'accuracy|repeatability|maximum measured error|low flow|0\.2%|0\.85|\b10\b.*m³',t,re.I):
                print('MATCH', name, i+1, re.sub(r'\s+', ' ', t)[:180])
    elif path.suffix == '.docx':
        with zipfile.ZipFile(path) as z:
            xml = etree.fromstring(z.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main', 'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        lines=[]
        for i,p in enumerate(xml.xpath('//w:body//w:p',namespaces=ns)):
            t=''.join(p.xpath('.//w:t/text() | .//m:t/text()',namespaces=ns))
            if t: lines.append(f'P{i+1}: {t}')
        (OUT/f'{name}.txt').write_text('\n'.join(lines),encoding='utf-8')
    else:
        wb = load_workbook(path, read_only=True, data_only=False)
        values = load_workbook(path, read_only=True, data_only=True)
        lines=[]
        for ws in wb:
            lines.append(f'=== SHEET {ws.title} ({ws.max_row}x{ws.max_column}) ===')
            for row, value_row in zip(ws, values[ws.title]):
                cells=[]
                for c, value_cell in zip(row, value_row):
                    if c.value is not None:
                        v=value_cell.value
                        cells.append(f'{c.coordinate}={c.value!s}' + (f' [cached={v}]' if c.data_type=='f' else ''))
                if cells: lines.append(' | '.join(cells))
        (OUT/f'{name}.txt').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
